"""Generate the Chinese analytical report from verified pipeline outputs only."""
from common import *
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def read(n):return pd.read_csv(TABLE/(n+'.csv'))
def jread(n):return json.loads((OUT/n).read_text(encoding='utf-8'))
def pct(x):return '不适用' if pd.isna(x) else f'{x:.1%}'
def pp(x):return '不适用' if pd.isna(x) else f'{100*x:+.2f}'
def mil(x):return f'{x/1e6:,.2f}'
def num(x):return f'{x:,.0f}'
def clean(s):return str(s).replace('UNKNOWN_SKU_','INN未知 SKU ')

def main():
    d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.72)
    for name in ['Normal','Body Text','Title','Heading 1','Heading 2','Caption']:
        st=d.styles[name];st.font.name='Microsoft YaHei';st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');st.font.color.rgb=RGBColor.from_string('17242D')
    d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.line_spacing=1.16;d.styles['Normal'].paragraph_format.space_after=Pt(7)
    d.styles['Title'].font.size=Pt(27);d.styles['Title'].font.color.rgb=RGBColor(0,0,0)
    d.styles['Heading 1'].font.size=Pt(19);d.styles['Heading 1'].paragraph_format.space_after=Pt(12)
    d.styles['Heading 2'].font.size=Pt(12);d.styles['Caption'].font.size=Pt(9)
    footer=sec.footer.paragraphs[0];footer.alignment=2;footer.add_run('药品组合分析  ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    md=[]
    def heading(title,first=False):
        if not first:d.add_page_break()
        d.add_heading(title,0 if first else 1);md.append('# '+title+'\n')
    def para(text):d.add_paragraph(text);md.append(text+'\n')
    def sub(text):d.add_heading(text,2);md.append('## '+text+'\n')
    def table(headers,rows,widths=None):
        tab=d.add_table(rows=1,cols=len(headers));tab.style='Light Shading Accent 1';tab.autofit=False
        for i,c in enumerate(tab.columns):c.width=Inches(widths[i] if widths else 7.06/len(headers))
        for i,h in enumerate(headers):tab.rows[0].cells[i].text=str(h)
        repeat=OxmlElement('w:tblHeader');tab.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells=tab.add_row().cells
            for c,v in zip(cells,row):c.text=clean(v)
        for row in tab.rows:
            cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
            for i,c in enumerate(row.cells):
                c.width=Inches(widths[i] if widths else 7.06/len(headers))
                for p in c.paragraphs:
                    p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
                    for r in p.runs:r.font.size=Pt(8.5)
        md.append('| '+' | '.join(map(str,headers))+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+'\n'.join('| '+' | '.join(clean(v).replace('|','/') for v in row)+' |' for row in rows)+'\n')
    def fig(name,caption,width=6.95):
        d.add_picture(str(OUT/'figures'/name),width=Inches(width));d.add_paragraph(caption,'Caption');md.append(f'![{caption}](../outputs/figures/{name})\n')
    summary=read('dataset_summary').iloc[0];annual=read('portfolio_annual');conc=read('portfolio_concentration').iloc[0];k=read('commercial_actions');known=k[k.known_inn]
    fs=jread('qa/forecast_summary.json');fl=jread('models/forecast_lock.json');ol=jread('models/opportunity_lock.json');fm=read('forecast_model_metrics');om=read('opportunity_model_metrics');orm=read('opportunity_regression_metrics');mapping=read('inn_market_mapping');rules=jread('qa/decision_rules.json')
    ft=fm[fm.split.eq('locked_holdout')&fm.model.eq(fl['selected_global'])].iloc[0];ot=om[om.split.eq('locked_test_2023H1')&om.model.eq(ol['classifier'])].iloc[0]
    rt=orm[orm.split.eq('locked_test_2023H1')&orm.model.eq(ol['regressor'])].iloc[0];base=orm[orm.split.eq('locked_test_2023H1')&orm.model.eq('Zero change')].iloc[0]
    heading('药品组合市场分析与需求预测',True)
    para('全量数据研究报告  2019—2023 年历史分析及 2024 年规划预测')
    para('AI-Driven Pharmaceutical Portfolio Analytics: Market Dynamics, Commercial Opportunity and Demand Forecasting')
    sub('主要结论')
    para(f'全量样本覆盖 {int(summary.sku_count):,} 个 SKU、{int(summary.inn_known_count):,} 个可识别 INN 和 {int(summary.atc4_count):,} 个 ATC4 类别。2023 年销量为 {mil(annual.sales.iloc[-1])} 百万原始单位，同比 {pct(annual.YoY.iloc[-1])}，较 2019 年变化 {pct(annual.sales.iloc[-1]/annual.sales.iloc[0]-1)}。组合总量接近企稳，但不能据此推断所有治疗领域或分子均已恢复。')
    para(f'按时间验证选定的需求预测模型为 {fl["selected_global"]}，2023 年冻结留出集 WAPE 为 {pct(ft.WAPE)}，Bias 为 {pct(ft.Bias)}。用截至 2023 年底的数据重拟合后，2024 年规划点预测为 {mil(fs["point_2024"])} 百万原始单位，相对 2023 年变化 {pct(fs["growth_2024"])}。该结果是历史时点预测，不是对当前年份的预测。')
    para(f'机会分类模型选定 {ol["classifier"]}，留出集按锚点平均 Lift@20% 为 {ot["Lift@20"]:.2f} 倍。连续份额变化模型选定 {ol["regressor"]}，留出 MAE 为 {rt.MAE*100:.2f} 个百分点；零变化基线为 {base.MAE*100:.2f} 个百分点。排序价值与份额增量预测精度应分别判断。')
    sub('建议的使用方式')
    para('商业团队以同类份额下降的大体量分子作为防守诊断名单，以份额上升且机会排名靠前的分子作为增长核查名单。供应团队单独审查预测增长、需求波动与历史低估同时出现的品种。行动规则只安排分析与复核优先级，不直接决定停产、利润投入或临床替代。')
    para('所有销量保留 CSV 原始数值与小数。案例按包装销量描述数据，但原文件没有独立的倍率、地区和渠道字典，因此本报告统一称“原始销量单位”。')

    heading('一 研究范围与数据证据')
    para('研究问题是如何结合历史销量、治疗层级与产品特征，识别需求变化、竞争位置和未来规划重点。分析包括组合描述、治疗同类比较、产品属性关联、机会排序、需求预测及透明行动规则。')
    table(['指标','全量结果'],[['SKU',num(summary.sku_count)],['可识别 INN',num(summary.inn_known_count)],['原始 INN 标签数',num(summary.inn_raw_labels)],['INN 未知 SKU',num(summary.unknown_inn_skus)],['品牌编码',num(summary.brand_count)],['ATC4 / ATC5',f'{int(summary.atc4_count)} / {int(summary.atc5_count)}'],['月份 / SKU 月观测',f'60 / {int(summary.sku_month_observations):,}'],['2019至2023累计销量',f'{mil(summary.total_sales_2019_2023)} 百万原始单位']],widths=[3,4.06])
    para(f'INN 占位符归一化后视为未知；每个未知 SKU 保留独立键，不把 {int(summary.unknown_inn_skus)} 个 SKU 合并成一个虚构分子。这部分占五年总销量的 {pct(summary.unknown_inn_sales_share)}，仍纳入组合总量与 SKU 预测。')
    para('数据源为 teva_sales.csv；上下文材料为 MLapps_Teva_Pharma.pdf。品牌字段为编码，不含生产商身份。样本覆盖率、全国市场边界、产品历史上下架名单均未给出，因而不能将样本组合解释为 Teva 公司全部销量，也不能将同类份额称为全国市场份额。')
    para('本次使用全量 13,762 个 SKU。过去的小样本数字和模型结果均未作为本次结论依据。源文件 SHA256 与运行环境版本保存在 QA 清单，可用于复核。')

    heading('二 数据质量与产品信息标准化')
    coverage=read('attribute_parse_coverage');ds=read('demand_segments')
    table(['检查','结果'],[['重复 SKU / 缺失销量 / 负销量','0 / 0 / 0'],['零销量观测',f'{int(summary.zero_sales):,}，占 {summary.zero_sales/summary.sku_month_observations:.1%}'],['含小数的销量观测',num(summary.fractional_sales_observations)],['出现内部零间隔的 SKU',num((ds.internal_zero_months>0).sum())],['末端连续零销量至少6个月的 SKU',num((ds.trailing_zeros>=6).sum())],['可比较单一质量剂量解析率',pct(coverage.loc[coverage.attribute.eq('strength_mg'),'coverage'].iloc[0])],['包装数量解析率',pct(coverage.loc[coverage.attribute.eq('pack_size'),'coverage'].iloc[0])]],widths=[4.1,2.96])
    para('数据工程将 60 个月宽表转换为 SKU × Month 长面板，并单独保存产品属性。剂量归一化为 mg；1 g 转为 1,000 mg，1 mcg 转为 0.001 mg。mg/ml 浓度、复方多剂量和容器重量不会强行合并为单一剂量。包装数量按 #20、pack of 30 等表达解析；无法可靠解析的值保留缺失标记。')
    para('前导零表示首次可观测销售前的空窗，不能确认实际上市日期；内部零间隔表示正销量之间的零销量月份，称为潜在可用性事件；末端零可能涉及退出、停供或需求减弱，统一作为生命周期核查信号。全零 SKU 单独分类。')
    para('原始零值不会被库存假设替换。由于缺少缺货标签、库存、处方或订单数据，预测对象是观测销量而非不受供给限制的潜在需求。质量标记用于分层与解释，不能证明真实断供。')

    heading('三 全组合需求变化')
    fig('01_portfolio_sales.png','图1 组合年度与月度销量，单位为百万原始销量单位。')
    table(['年份','销量 百万','同比'],[[int(r.year),mil(r.sales),pct(r.YoY)] for r in annual.itertuples()],widths=[1.3,3,2.76])
    para(f'2019—2023 年组合复合增长率为 {pct(conc.portfolio_CAGR)}。2022 年同比变化为 {pct(annual.loc[annual.year.eq(2022),"YoY"].iloc[0])}，是年度总量变化的主要转折；2023 年的小幅回升尚不足以恢复此前规模。')
    para('月度序列存在重复波动及明显年度水平变化。图中转折仅为描述性观察，未使用外部事件、政策或疫情变量识别因果，也未将趋势断点等同于某项外部事件的影响。')

    heading('四 治疗领域贡献与市场定义')
    fig('02_atc_landscape.png','图2 ATC1 治疗领域贡献，按2023年销量排序。')
    para(f'2023 年已知 INN 前10名和前20名分别占完整组合的 {pct(conc.Top10_known_INN_share_of_portfolio)} 与 {pct(conc.Top20_known_INN_share_of_portfolio)}。在已知 INN 销量内重新归一化后，HHI 为 {conc.HHI_known_INN_normalized:.4f}，即 {conc.HHI_known_INN_normalized*10000:.1f} 点。此处反映分子销量集中度，不用于监管意义上的市场集中度认定。')
    counts=mapping.market_level.value_counts()
    table(['分子竞争映射','INN 数'],[['ATC4 同类市场',int(counts.get('ATC4',0))],['ATC3 回退',int(counts.get('ATC3',0))],['不报告分子级竞争份额',int(counts.get('Excluded',0))]],widths=[4.8,2.26])
    para('优先使用唯一 ATC4；同类可识别 INN 不足3个时回退 ATC3。跨多个 ATC4 但全部属于同一 ATC3 的分子，在 ATC3 比较；跨多个 ATC3 的分子不输出单一竞争份额。历史机会锚点进一步按当时已有正销量的 INN 重新检查最低同类要求。')
    para('分母包括相应类别内全部样本 SKU，含 INN 未知项。年度报告份额按全年销量相除；机会预测目标使用月末对应月份的销量份额。两者时间口径不同，结果表已分别命名。')

    heading('五 市场增长与竞争位置分解')
    fig('03_market_vs_share.png','图3 样本市场增长与年度份额变化，气泡大小表示分子销量规模。')
    de=read('growth_decomposition');top=de.nlargest(5,'share_effect')
    table(['INN','市场效应 百万','份额效应 百万','份额变化 百分点'],[[r.INN,mil(r.market_effect),mil(r.share_effect),pp(r.share_change_12m)] for r in top.itertuples()],widths=[2.55,1.5,1.5,1.51])
    para('对每个可比较分子，将销量表示为份额与样本市场规模的乘积。2022到2023年的变化分为三项：原份额乘市场规模变化、原市场规模乘份额变化，以及两种变化的交互项。三项之和与实际销量变化逐行核对。')
    para('市场扩大但分子份额下降，提示需要诊断相对竞争位置；市场稳定或下降时仍获得份额，才更接近相对增长信号。销量与份额同时下降的分子进入风险核查。上述标签是观测结果分类，不能确定竞争营销、价格或疗效是原因。')

    heading('六 需求形态与生命周期')
    seg=read('demand_segment_summary')
    table(['需求形态','SKU 数','平均零值率','平均 CV'],[[r.segment,int(r.sku_count),pct(r.mean_zero_rate),f'{r.mean_cv:.2f}'] for r in seg.itertuples()],widths=[3.1,1.2,1.35,1.41])
    para('需求特征在 SKU 和 INN 两级分别计算，包括零值率、均值、CV、正销量大小的 CV²、平均需求间隔 ADI、季节性强度、趋势强度和12阶自相关。ADI 等于观测月份数除以正销量月份数；全零序列不把 ADI 填成正常数值。')
    para('规则先识别全零、末端连续6个月零销量、最近12个月才出现首次正销量等生命周期模式，再按 ADI 1.32 和正销量 CV² 0.49 区分间歇与块状波动；其余使用季节性强度0.5和 CV 1作为描述性分界。阈值是建模规则，不能视为普遍临床或业务标准。')
    para('季节性与趋势强度基于 log1p 销量的线性趋势和月份虚拟变量分解，并用残差方差构造0到1指标；本实现没有将该指标冒称为 STL。预测回测时仅用起点前历史重新计算分层，避免拿2023年的生命周期结果指导2022年底模型选择。')
    para('稳定需求通常更适合平滑基线；间歇需求需要对发生概率和非零规模分别处理。具有末端零或极短历史的品种，数值预测仅供人工复核，不能自动解释为停产。')

    heading('七 产品属性与销量的关联')
    attr=read('attribute_fixed_effects');meta=jread('qa/attribute_model.json');ml=read('attribute_ml_metrics');imp=read('attribute_permutation_importance')
    para(f'解释模型使用 {meta["complete_case_skus"]:,} 个剂量和包装均可解析、INN 已知的 SKU，共 {meta["observations"]:,} 条月度观测。因变量为 log1p 销量，解释变量为 log1p 包装数、log1p 剂量和剂型；吸收 INN 与日历月份的交互固定效应，标准误按 SKU 聚类。')
    table(['变量','系数','95%区间'],[[r.term,f'{r.coefficient:+.3f}',f'[{r.ci95_low:+.3f}, {r.ci95_high:+.3f}]'] for r in attr.itertuples()],widths=[3.65,1.1,2.31])
    pack=attr[attr.term.eq('log1p_pack_size')].iloc[0];strength=attr[attr.term.eq('log1p_strength_mg')].iloc[0]
    para(f'包装数系数为 {pack.coefficient:+.3f}，剂量强度系数为 {strength.coefficient:+.3f}。剂量区间是否包含零：{"是" if strength.ci95_low<=0<=strength.ci95_high else "否"}。包装数与包装销量的关系包含计量机制：大包装每盒可提供更多用量，销量盒数较低并不等于治疗用量或患者偏好较低。')
    para('在同一分子与同一月份内比较不同 SKU，以避免把不同分子的规模与时间变化当成属性效应。SKU 固定效应未加入，因为包装、剂量等时间不变属性会被其完全吸收。所得系数属于条件相关关系；价格、渠道和厂家缺失，不能解释为改变包装或剂量会造成相同幅度销量变化。')
    para('补充模型为全局梯度提升树，使用产品属性、分子/ATC编码、月份、已知滞后销量和滚动均值。训练截至2021年，2022年用于置换重要度；2023年仅作一步预测评估。每个月允许使用此前已经发生的实际销量，因此其误差不可与一次性12个月预测直接比较。')
    table(['2023 一步预测','结果'],[['WAPE',pct(ml.loc[ml.split.eq('holdout_2023'),'WAPE'].iloc[0])],['置换重要度前3项','、'.join(imp.head(3).feature)]],widths=[2.4,4.66])

    heading('八 商业组合分群')
    cluster=jread('qa/cluster_summary.json');fig('04_portfolio_pca.png','图4 标准化特征的 PCA 投影，颜色表示描述性聚类。')
    para(f'K-means 比较 K=3到6，以轮廓系数选择 K={cluster["selected_K"]}。输入包含规模、增长、同类份额变化、市场增长、波动、季节性、零值率、SKU与品牌及属性覆盖、预测增长。对极端比例按1%和99%分位截尾、缺失值中位数填补，再做 Z 标准化。')
    c=known.groupby(['cluster','archetype']).agg(n=('INN','size'),sales=('sales_2023','sum'),growth=('YoY','median')).reset_index()
    table(['群组','描述','INN 数','销量占比'],[[f'C{int(r.cluster)}',r.archetype,r.n,pct(r.sales/known.sales_2023.sum())] for r in c.itertuples()],widths=[.65,3.5,1,1.91])
    para(f'前两个主成分解释标准化特征方差的 {sum(cluster["PCA_variance_explained"]):.1%}，只用于展示，不参与聚类或行动规则。群组名称依据中位特征描述，不强行制造六种商业原型。同一原型可包含多个群组，聚类结果不直接等于投资建议。')

    heading('九 未来商业机会的定义与验证')
    para('观察单位为 INN × 锚点月。主目标是6个月后单月同类份额减锚点月份额；辅助目标为份额增量为正、且在同锚点合格分子中位列前25%的未来赢家。边界并列按 INN 排序，保证标签可复现。')
    para('特征只使用锚点月结束前可观测的信息：销量窗口、历史增长、同类份额及动量、市场增长、相对增长、波动、季节相关性、零值率、ADI、已出现销售的产品结构及生命周期。静态治疗分类来自文件提供的分类快照，历史名录完整性仍是限制。')
    table(['阶段','锚点与标签边界'],[['开发训练','2020年1月至12月锚点；最晚标签为2021年6月'],['验证选型','2021年7月至12月锚点；标签落在2022年上半年'],['测试前重拟合','使用最晚2022年6月锚点；标签截至2022年12月'],['冻结留出','2023年1月至6月锚点；标签在2023年7月至12月'],['最终规划','截至2023年12月已知标签重拟合，预测2024年6月份额变化']],widths=[1.65,5.41])
    para('训练与验证之间保留6个月标签隔离，防止未来份额变化跨入下一阶段。分类以逐锚点平均 Lift@20%选型，PR-AUC用于辅助；回归以验证 MAE选型。模型池包含基线、正则化线性模型、随机森林和梯度提升。未来同类市场销量为零时，份额目标不可定义，该观察不参与监督评价，不将其强行填成零。')
    para('Precision@10%与@20%、Lift、NDCG逐锚点排序后平均；PR-AUC、ROC-AUC和 Brier 使用合并留出观测。校准表展示预测概率与实际频率，但没有把未校准分类输出宣称为可靠的绝对发生概率。')

    heading('十 商业机会模型的结果')
    fig('05_opportunity_lift.png','图5 机会排名的验证与留出表现，随机基准的 Lift 约为1。')
    t=om[om.split.eq('locked_test_2023H1')]
    table(['模型','PR-AUC','P@20%','Lift@20%','Brier'],[[r.model,f'{r.PR_AUC:.3f}',pct(getattr(r,'_7',np.nan)), '', ''] for r in []]) if False else None
    table(['模型','PR-AUC','P@20%','Lift@20%','Brier'],[[r['model'],f'{r["PR_AUC"]:.3f}',pct(r['Precision@20']),f'{r["Lift@20"]:.2f}',f'{r["Brier"]:.3f}'] for _,r in t.iterrows()],widths=[2.2,1.05,1.3,1.3,1.21])
    failed=om.loc[(om.split.eq('validation_2021H2'))&(~om.converged),'model'].tolist() if 'converged' in om else []
    if failed:para('数值诊断：'+ '、'.join(failed)+' 在验证拟合中达到迭代上限，结果仅作诊断参考，不参与最终选型。收敛状态已逐模型写入指标表。')
    para(f'锁定模型 {ol["classifier"]} 的前20%名单精确率为 {pct(ot["Precision@20"])}，测试基础赢家率为 {pct(ot.base_rate)}。Lift 大于1表示在该历史样本中优于随机挑选，并不表示新增商业投入会获得同等倍数回报。')
    para(f'主回归目标的最佳模型 {ol["regressor"]} 在留出集 MAE 为 {rt.MAE*100:.2f} 个百分点，相比零变化基线 {base.MAE*100:.2f} 个百分点。若基线更好，应承认连续增量预测尚无稳定优势，同时单独检查排序任务是否仍有筛选价值。')
    unc=jread('qa/opportunity_uncertainty.json');para(f'按6个测试锚点重采样得到的 Lift@20% 描述性95%范围为 {unc["anchor_bootstrap_95_low"]:.2f}—{unc["anchor_bootstrap_95_high"]:.2f}。未来窗口彼此重叠，因此这不是独立样本意义上的显著性证据；可用测试年份不足以支撑稳定泛化承诺。')

    heading('十一 十二个月销量预测')
    fig('06_forecast_comparison.png','图6 验证与冻结留出集 WAPE。最终模型由验证结果选定。')
    para('模型包括季节性朴素法、加性季节调整后的阻尼 ETS、Ridge ARX 和全局梯度提升树。全局模型使用标准化销量滞后、滚动均值和标准差、零值率、ADI、月份、INN、主导ATC3/ATC4与剂型及剂量包装特征。高基数类别按起点销量保留前254类，其余合并；岭回归独热编码，提升树使用原生类别分裂。预测递归推进，不注入未来实际销量。')
    para('验证的12个月起点为2020年12月、2021年6月和2021年12月；2022年3月与6月增加6个月检验，2022年9月增加3个月检验。所有验证目标均截至2022年底。12个月选型按三个起点 WAPE的平均值，模型锁定后才计算2023年排名。')
    t=fm[fm.split.eq('locked_holdout')]
    table(['模型','WAPE','Bias','MASE'],[[r.model,pct(r.WAPE),pct(r.Bias),f'{r.MASE:.2f}'] for r in t.itertuples()],widths=[3.45,1.2,1.2,1.21])
    para('WAPE按分子月绝对误差汇总，不允许不同分子的高估与低估互相抵销；Bias保留方向；MASE以各序列训练期12阶季节差分绝对值均值缩放，分母为零的序列不计入 MASE，并单独记录数量。另提供 RMSE 与按需求形态选模的敏感性结果。')

    heading('十二 不确定性与层级一致性')
    coverage=read('forecast_interval_coverage');hm=read('hierarchical_forecast_metrics');im=read('intermittent_forecast_metrics');sl=jread('models/sku_forecast_lock.json')
    para('年度上下界使用验证期完整12个月路径的年度误差校准，月度上下界使用对应预测步长误差；按销量规模分层并将下界截断为零。P50_proxy 是规划点预测，不是严格估计的条件中位数。P10和P90为经验残差规划范围，有限且重叠的起点使其概率解释较弱。')
    table(['规模组','序列数','2023年度80%范围覆盖率'],[[int(r.scale_group),int(r.n_series),pct(r.annual_80_coverage)] for r in coverage.itertuples()],widths=[1.4,1.5,4.16])
    para('不能把各月或各分子的分位数直接相加并称为组合分位数。组合年度范围另用组合年度误差构造，只有三个验证路径，应按情景范围使用；点预测可以正常加总。范围覆盖率用于揭示不确定性估计的不足，不用2023年结果回头修改已锁定方法。')
    para(f'三个规模组的实际年度覆盖率为 {coverage.annual_80_coverage.min():.1%} 至 {coverage.annual_80_coverage.max():.1%}，均低于名义80%。经验范围的校准不足，业务使用时应保留额外的人工情景审查，不应把P10/P90当作已充分校准的承诺。')
    table(['层级方案','2023 WAPE','Bias'],[[r.method,pct(r.WAPE),pct(r.Bias)] for r in hm.itertuples()],widths=[3.5,1.78,1.78])
    para(f'底层 SKU 的间歇需求比较 Croston、SBA、TSB 与季节朴素法，验证选定 {sl["intermittent_model"]}；非间歇 SKU 使用季节朴素法。底层预测汇总到 INN，与直接 INN 预测进行同口径比较。')
    para('最终规划仍保留预先选定的直接 INN 模型，再按最近12个月 SKU 销量结构分配，分别汇总到 ATC5、ATC4 与组合。这保证加总一致，但固定产品结构假设可能遗漏新品或产品迁移；分配结果不能当作独立优化过的 SKU 预测。')

    heading('十三 潜在供给空窗与替代线索')
    submeta=jread('qa/substitution_summary.json');ev=read('substitution_inn_summary').sort_values('events',ascending=False)
    para(f'同一可识别 INN 内检出 {submeta["events"]:,} 个内部零销量事件，覆盖 {submeta["eligible_inns"]:,} 个分子。以事件前3个月同分子其他 SKU 的月均销量为基线，{pct(submeta["positive_peer_uplift_share"])} 的有效事件出现其他 SKU 销量上升；中位变化为 {pct(submeta["median_peer_uplift"])}。')
    table(['INN','内部零事件','正向同类提升占比','中位提升'],[[r.INN,int(r.events),pct(r.positive_uplift_rate),pct(r.median_peer_uplift)] for r in ev.head(8).itertuples()],widths=[2.9,1.2,1.65,1.31])
    para('事件要求零销量前后均有正销量，不把前导零和末端零混作短期可用性问题。输出包含零间隔起止、时长、焦点 SKU 事前规模、其他 SKU 事前与期间规模、同比同月份基线、表观吸收比例。')
    para('另对每个分子规模前8个 SKU 的 log1p 销量一阶差分计算相关性，保留低于−0.2的探索性边，并记录剂型及剂量。该列表可支持后续产品关系可视化；未训练图神经网络。')
    para('这些结果没有未受影响的对照组，也未剔除全部季节、促销和共同冲击。其他 SKU 的同期增长只是待验证的替代线索，不能认定因果性蚕食或已确认缺货。同分子不同给药途径、剂型和剂量也不必然能够临床互换。')
    para('建议下一步将重点事件与库存、发货、停供和渠道记录连接，再按同剂型、可比剂量及包装建立更严格的对照事件研究。')

    heading('十四 从模型到商业行动')
    fig('07_growth_and_actions.png','图7 增长预期与份额动量。行动由公开阈值生成，聚类不参与决策。')
    act=read('commercial_action_summary')
    table(['行动','分子或未知SKU数','2023销量占比'],[[r.action,int(r.inn_or_unknown_sku_count),pct(r.sales_2023/annual.sales.iloc[-1])] for r in act.itertuples()],widths=[3.2,1.9,1.96])
    para('Invest：规模前25%、份额上升、预测增长为正且机会概率前20%；Emerging Opportunity：规模尚未进入前25%、份额上升且机会概率前20%，先验证增长驱动；Defend：规模前25%但份额下降；At Risk：份额与预测销量同时下降。市场带动增长定义为销量增长且份额变化不超过0.2个百分点。')
    para(f'Supply Watch 是可叠加标签：预测增长超过10%，同时 CV>1或零值率>20%，且历史偏差低于−10%。本次有 {int(k.supply_watch.sum())} 个序列符合。未知 INN 进入 Data Review；竞争份额不适用的分子进入 Monitor。')

    heading('十五 2024 年规划表与优先核查名单')
    para(f'组合点预测为 {mil(fs["point_2024"])} 百万原始单位，增长 {pct(fs["growth_2024"])}。组合经验情景范围为 {mil(fs["P10_scenario"])} 至 {mil(fs["P90_scenario"])} 百万，基于三个历史年度误差路径，不能视为经过充分验证的80%概率区间。')
    table(['INN','2023实际 百万','2024点预测 百万','增长','行动'],[[r.INN,mil(r.sales_2023),mil(r.P50_proxy),pct(r.growth_2024),r.action] for r in known.nlargest(10,'sales_2023').itertuples()],widths=[2.65,1.15,1.2,.86,1.2])
    sub('重点名单的使用顺序')
    defend=known[known.action.eq('Defend')].nlargest(3,'sales_2023');invest=known[known.action.isin(['Invest','Emerging Opportunity'])].nlargest(3,'opportunity_probability')
    para('优先防守诊断：'+('、'.join(defend.INN) if len(defend) else '暂无符合规则的分子')+'。首先核查同治疗类别竞争品、渠道覆盖与供给变化，区分样本市场下行和自身份额流失。')
    para('增长驱动核查：'+('、'.join(invest.INN) if len(invest) else '暂无符合规则的分子')+'。结合价格、营销触达、订单与渠道数据后，再讨论增量资源配置。')
    para('完整规划表包含每个 INN 的2023实际、2024点预测、年度经验上下界、预测增长、年度份额动量、机会模型输出、行动与 Supply Watch。所有未知 INN 的销量与预测也保留在表中，但不作为分子商业推荐。')

    heading('十六 治疗领域深入分析')
    deep=jread('qa/deep_dive.json');dt=read('therapeutic_deep_dive')
    para(f'选取 {deep["ATC4"]}。该类别包含 {deep["known_inn_count"]} 个已知 INN。筛选要求4至12个分子、至少4个分子各占1%以上份额、最大分子份额低于70%，再选择2023销量最大的类别。市场规模为 {mil(deep["sales_2023"])} 百万原始单位，同比变化 {pct(deep["market_growth_2023"])}。')
    fig('08_therapeutic_deep_dive.png','图8 市场规模、主要分子份额、增长分解及2024分配预测。',6.8)
    winners=dt[dt.share_change>0].sort_values('share_change',ascending=False).head(2);losers=dt[dt.share_change<0].sort_values('share_change').head(2)
    para('份额提升较明显：'+'；'.join(f'{r.INN} {pp(r.share_change)} 个百分点' for r in winners.itertuples())+'。份额下滑较明显：'+'；'.join(f'{r.INN} {pp(r.share_change)} 个百分点' for r in losers.itertuples())+'。')
    para('此处竞争分析使用 INN × ATC4 切片，因此跨多个治疗类别的分子也可以在本类别内比较。2024预测由 SKU 分配后汇总，仅覆盖该治疗类别，避免把分子在其他适应类别的销量算入本市场。')

    heading('十七 解释边界与后续验证')
    sub('本研究能支持的结论')
    para('可以回答样本组合的销量变化、治疗类别贡献、同类份额动量、需求形态，以及在规定历史回测中的模型表现。可将分子与 SKU 排入商业诊断、供给复核和数据完善名单。增长分解是代数恒等式，产品属性、替代关系和模型重要度均为相关性分析。')
    sub('必须保留的限制')
    para('一是市场代表性与数据单位。缺少地区、渠道、样本覆盖率及单位倍率字典，同类份额仅在样本内成立。二是名录与历史分类。提供的是固定 SKU 和治疗分类快照，无法完整重建历史退市、变更及未纳入产品，可能存在幸存者偏差。')
    para('三是销量不等于潜在需求。零销量可能反映需求、供给、产品生命周期或数据记录。没有真实 OOS 标签，无法确认缺货。四是属性解析和混杂。完整病例筛选会改变样本构成；同一 INN 不同剂型的剂量仍未必可比较，生产商、价格和渠道缺失限制解释。')
    para('五是预测和分类的外推。只有五年历史，2023只提供一个冻结年份；短锚点和重叠未来窗口增加评价不确定性。树模型的类别编码只提供预测性编码，不能把其数值顺序视为治疗距离。六是区间校准。只有三个年度误差路径，不具备稳定尾部覆盖证据。')
    para('七是经济决策范围。没有价格、成本、毛利、商业活动或战略重要性，因此没有计算利润、投资回报或停产建议。没有临床可替代性评估，不将观察到的同类提升用于患者治疗选择。')
    sub('建议新增的数据')
    para('优先连接库存与供货事件、价格及渠道覆盖、厂家与促销投入、地区和外部市场规模、产品真实上市/退出时间。取得2024及以后实际销量后，应按预先锁定的代码进行新年度外部检验，而不是重新选择历史上最好看的模型。')

    heading('十八 复现方法与交付索引')
    para('分析代码按数据审计、特征、组合、市场、需求形态、属性、机会、预测、层级、替代与决策拆分。所有 CSV 采用 UTF-8 BOM，可直接用 Excel 打开；长面板与中间矩阵使用 Parquet 和 NPY 保存。模型文件、选择记录与 QA 输出保留在项目中。')
    table(['内容','位置'],[['一键运行','run_pipeline.py'],['模块化代码','analysis/01 至 13，各模块职责见 README'],['十张核心表与辅助结果','outputs/tables'],['八张核心图','outputs/figures'],['模型与选型记录','outputs/models'],['源文件指纹与核验','outputs/qa'],['中文报告','report/中文分析报告.docx 和 .md']],widths=[2.3,4.76])
    para('建议复现顺序：安装 requirements.txt；运行 python run_pipeline.py；运行 python analysis/verify_pipeline.py。也可用 --from-stage 和 --to-stage 从指定阶段重新执行。每个预测模型的选型配置先落盘，再计算留出结果，防止按2023年排名事后换模。')
    sub('资料与方法来源')
    para('原始数据：teva_sales.csv，2019年1月至2023年12月。背景：Laura Tolosi-Halacheva，Machine Learning applications at Teva Pharma with focus on forecasting，2024年3月，重点参考第17至24页及第28至31页。')
    para('预测验证和层级一致性参考 Hyndman 与 Athanasopoulos 的 Forecasting: Principles and Practice 第三版：https://otexts.com/fpp3/tscv.html 及 https://otexts.com/fpp3/reconciliation.html。')
    para('置换重要度和轮廓系数实现参考 scikit-learn 官方文档：https://scikit-learn.org/stable/modules/permutation_importance.html 及 https://scikit-learn.org/stable/modules/generated/sklearn.metrics.silhouette_score.html。软件版本记录见 requirements.txt 和 QA 运行清单。')
    d.core_properties.title='药品组合市场分析与需求预测';d.core_properties.subject='2019至2023全量分析与2024历史规划预测';d.core_properties.author=''
    for root in [d.element,d.styles.element]:
        for border in list(root.iter(qn('w:pBdr'))):border.getparent().remove(border)
    dest=ROOT/'report/中文分析报告.docx';d.save(dest);(ROOT/'report/中文分析报告.md').write_text('\n'.join(md),encoding='utf-8')
    jsave({'document':str(dest),'paragraphs':len(d.paragraphs),'tables':len(d.tables),'figures':len(d.inline_shapes),'sections':19},OUT/'qa/document_structure.json')
    print('Report created',dest,flush=True)
if __name__=='__main__':main()
