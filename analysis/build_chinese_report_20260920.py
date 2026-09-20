"""Chinese translation with explicit audit annotations; original files are retained."""
from pathlib import Path
import json, math
import numpy as np
import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/'outputs/iqvia_teva_case_20260919'
QA=BASE/'verification_20260920'
FIG=QA/'figures_zh'; FIG.mkdir(exist_ok=True,parents=True)
R=json.loads((QA/'recomputed/analysis_results.json').read_text(encoding='utf-8'))
V=json.loads((QA/'verification_summary.json').read_text(encoding='utf-8'))
K=pd.DataFrame(R['inn_kpi'])
M={'ACIDUM ACETYLSALICYLICUM*':'乙酰水杨酸','IBUPROFENUM':'布洛芬','PARACETAMOLUM':'对乙酰氨基酚','DICLOFENACUM':'双氯芬酸','BISOPROLOLUM':'比索洛尔','METFORMINUM':'二甲双胍','CEFTRIAXONUM':'头孢曲松','AMLODIPINUM':'氨氯地平','FLUCONAZOLUM':'氟康唑','PANTOPRAZOLUM':'泮托拉唑','AZITHROMYCINUM':'阿奇霉素','AMBROXOLUM':'氨溴索','ROSUVASTATINUM':'瑞舒伐他汀','MELOXICAMUM':'美洛昔康','PREGABALINUM':'普瑞巴林','SILDENAFILUM':'西地那非','ACETYLCYSTEINUM':'乙酰半胱氨酸','ATORVASTATINUM':'阿托伐他汀','LEVOFLOXACINUM':'左氧氟沙星','CITICOLINUM':'胞磷胆碱'}
MODEL={'Logistic L2':'L2 逻辑回归','Elastic Net':'弹性网逻辑回归','Random Forest':'随机森林','Boosted Trees':'提升树','Seasonal Naive':'季节性朴素法','Damped ETS':'阻尼 ETS','Global Ridge ARX':'全局岭回归 ARX','Global Boosted Trees':'全局提升树','Validation-weighted Ensemble':'验证加权集成'}
CL={'Growth / Emerging':'成长与新兴型','Niche / Volatile':'利基与波动型','Mature Leaders':'成熟领先型','At-Risk / Declining':'风险与下滑型'}
ACT={'Growth Leader':'增长领先','Emerging Opportunity':'新兴机会','Defend':'防守','At Risk':'风险','Market-led Growth':'市场带动增长','Monitor':'监测'}
pct=lambda x:f'{x:.1%}'
million=lambda x:f'{x/1e6:.1f}'
pp=lambda x:f'{x*100:+.1f}'

def font(n,bold=False): return ImageFont.truetype('C:/Windows/Fonts/msyhbd.ttc' if bold else 'C:/Windows/Fonts/msyh.ttc',n)
colors=['#1C8A87','#2F6B8A','#D5A021','#B54A4A','#687780']
def canvas(title,sub='',height=760):
    im=Image.new('RGB',(1600,height),'white'); d=ImageDraw.Draw(im)
    d.text((60,25),title,font=font(35,True),fill='#15324B')
    d.text((60,80),sub,font=font(23),fill='#687780')
    return im,d
def plot_charts():
    a=R['portfolio_annual']; im,d=canvas('组合销量于 2021 年见顶后回落','单位 百万包装当量')
    for t in [0,40,80,120,160]:
        y=620-t/170*450; d.line((140,y,1510,y),fill='#D9E1E5',width=2); d.text((65,y-14),str(t),font=font(24),fill='#687780')
    for i,r in enumerate(a):
        x=210+i*270; y=620-r['sales']/1e6/170*450
        d.rectangle((x,y,x+130,620),fill=colors[0] if i<3 else colors[1]); d.text((x+20,640),str(r['year']),font=font(26),fill='#15324B'); d.text((x+20,y-42),million(r['sales']),font=font(27,True),fill='#15324B')
    im.save(FIG/'annual.png')
    im,d=canvas('分类测试期前 20% 排名提升倍数','按 2022 年 7 月至 2023 年 6 月全部观测合并排名',650)
    for i,r in enumerate(R['classification_results']):
        y=170+i*98; x=430+r['test_lift_top20']/3*970
        d.text((70,y),MODEL[r['model']],font=font(28),fill='#15324B'); d.rectangle((430,y,x,y+42),fill=colors[i]); d.text((x+20,y),f"{r['test_lift_top20']:.2f} 倍",font=font(28,True),fill='#15324B')
    im.save(FIG/'lift.png')
    f=R['forecast_2023_portfolio']; im,d=canvas('2023 年组合月度需求 实际值与预测值','组合汇总曲线 单位 百万包装当量')
    series=[('实际值',f['actual'],'#15324B'),('阻尼 ETS',f['Damped ETS'],colors[0]),('季节性朴素法',f['Seasonal Naive'],colors[2])]
    lo=4; hi=16
    for t in range(4,17,2):
        y=600-(t-lo)/(hi-lo)*420; d.line((150,y,1510,y),fill='#D9E1E5',width=2); d.text((80,y-15),str(t),font=font(24),fill='#687780')
    for name,vals,col in series:
        pts=[(150+i*1360/11,600-(v/1e6-lo)/(hi-lo)*420) for i,v in enumerate(vals)]; d.line(pts,fill=col,width=5)
    for i in range(12): d.text((135+i*1360/11,615),str(i+1),font=font(24),fill='#687780')
    for i,(name,_,col) in enumerate(series):
        x=220+i*430; d.line((x,710,x+50,710),fill=col,width=6); d.text((x+65,690),name,font=font(26),fill='#15324B')
    im.save(FIG/'forecast.png')
    im,d=canvas('四类组合特征的主成分投影','PCA 仅用于展示 聚类输入为标准化后的 14 项特征')
    xmin,xmax=K.pc1.min()-1,K.pc1.max()+1; ymin,ymax=K.pc2.min()-1,K.pc2.max()+1
    px=lambda x:160+(x-xmin)/(xmax-xmin)*1290
    py=lambda y:580-(y-ymin)/(ymax-ymin)*410
    for x in range(math.floor(xmin),math.ceil(xmax)+1,2): d.line((px(x),150,px(x),590),fill='#E0E6EA',width=2)
    for y in range(math.floor(ymin),math.ceil(ymax)+1,2): d.line((160,py(y),1450,py(y)),fill='#E0E6EA',width=2)
    labs=list(CL); used=[]
    for _,r in K.iterrows():
        x,y=px(r.pc1),py(r.pc2); d.ellipse((x-11,y-11,x+11,y+11),fill=colors[labs.index(r.cluster_label)])
        if r.INN in set(K.nlargest(7,'sales_2023').INN)|{'PREGABALINUM','ROSUVASTATINUM'}:
            ty=y-35
            while any(abs(ty-u[1])<25 and abs(x-u[0])<160 for u in used):ty+=30
            tx=min(x+18,1310); used.append((tx,ty)); d.text((tx,ty),M[r.INN],font=font(22),fill='#15324B')
    d.text((1400,608),'PC1',font=font(23),fill='#687780'); d.text((70,150),'PC2',font=font(23),fill='#687780')
    for i,lab in enumerate(labs):
        x=105+i*380; d.ellipse((x,696,x+20,716),fill=colors[i]); d.text((x+33,687),CL[lab],font=font(25),fill='#15324B')
    im.save(FIG/'clusters.png')
    im,d=canvas('2024 年预测增长与 2023 年组合份额动量','气泡面积近似表示 2023 年销量 份额变化单位为百分点')
    xmin,xmax=-.85,.30; ymin,ymax=-1.6,1.3
    px=lambda x:170+(x-xmin)/(xmax-xmin)*1280
    py=lambda y:590-(y-ymin)/(ymax-ymin)*410
    for x in [-.8,-.6,-.4,-.2,0,.2]:
        d.line((px(x),155,px(x),590),fill='#D9E1E5',width=2);d.text((px(x)-25,606),f'{x:.0%}',font=font(23),fill='#687780')
    for y in [-1.5,-1,-.5,0,.5,1]:
        d.line((170,py(y),1450,py(y)),fill='#D9E1E5',width=2);d.text((90,py(y)-15),f'{y:g}',font=font(23),fill='#687780')
    d.line((px(0),155,px(0),590),fill='#15324B',width=3);d.line((170,py(0),1450,py(0)),fill='#15324B',width=3)
    offsets={'IBUPROFENUM':(-40,-65),'PREGABALINUM':(-65,-55),'ACETYLCYSTEINUM':(-160,16),'PARACETAMOLUM':(-40,40),'ACIDUM ACETYLSALICYLICUM*':(-220,20),'LEVOFLOXACINUM':(-20,20)}
    for _,r in K.sort_values('sales_2023').iterrows():
        x,y=px(r.forecast_growth_2024),py(r.portfolio_share_change_2023*100); rad=9+30*math.sqrt(r.sales_2023/K.sales_2023.max())
        col=colors[0] if r.forecast_growth_2024>0 else colors[3];d.ellipse((x-rad,y-rad,x+rad,y+rad),fill=col,outline='white',width=2)
        if r.INN in offsets:
            dx,dy=offsets[r.INN]; d.text((x+dx,y+dy),M[r.INN],font=font(24,True),fill='#15324B')
    d.text((180,691),'横轴 预测同比增长率     纵轴 组合份额变化',font=font(26),fill='#687780')
    im.save(FIG/'matrix.png')

DOC=Document()
sec=DOC.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(1.9);sec.right_margin=Cm(1.9)
sec.footer_distance=Cm(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    s=DOC.styles[name];s.font.name='Microsoft YaHei';s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');s.font.color.rgb=RGBColor(0,0,0)
    s.font.size=Pt({'Normal':10.5,'Title':24,'Subtitle':12,'Heading 1':17,'Heading 2':12,'Caption':9}[name]);s.paragraph_format.space_after=Pt(5);s.paragraph_format.line_spacing=1.10
    for attr in ['w:asciiTheme','w:hAnsiTheme','w:eastAsiaTheme','w:cstheme']:
        s._element.get_or_add_rPr().rFonts.attrib.pop(qn(attr),None)
DOC.styles['Subtitle'].font.italic=False
for el in list(DOC.styles.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
DOC.styles['Normal'].paragraph_format.widow_control=True
for s in ['Heading 1','Heading 2']: DOC.styles[s].paragraph_format.keep_with_next=True
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def para(t,style=None): return DOC.add_paragraph(t,style)
def head(t,l=1): return DOC.add_heading(t,l)
def page(t):DOC.add_page_break();head(t)
def note(t):
    p=para(t);p.paragraph_format.space_after=Pt(7)
    for run in p.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('596773')
    return p
def table(headers,rows,widths=None,size=9):
    t=DOC.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    widths=widths or [17.2/len(headers)]*len(headers)
    for i,w in enumerate(widths):t.columns[i].width=Cm(w)
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h)
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row):cells[i].text=str(x)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for tag,n in [('top',75),('bottom',75),('start',90),('end',90)]:
                e=OxmlElement('w:'+tag);e.set(qn('w:w'),str(n));e.set(qn('w:type'),'dxa');mar.append(e)
            tcpr.append(mar)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EEF5' if ri==0 else 'FFFFFF');tcpr.append(sh)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.1
                if ri==0 or ci>0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:run.font.size=Pt(size);run.bold=(ri==0)
    spacer=DOC.add_paragraph();spacer.paragraph_format.space_after=Pt(0);spacer.paragraph_format.line_spacing=Pt(3);spacer.add_run(' ').font.size=Pt(3)
    return t
def figure(name,caption,width=16.4):
    p=DOC.add_paragraph();p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(2);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    s=p.add_run().add_picture(str(FIG/name),width=Cm(width));s._inline.docPr.set('descr',caption)
    p=para(caption,'Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
def location(code,label,path,extra=''):
    p=para(f'{code}  {label}');p.paragraph_format.keep_with_next=True
    for r in p.runs:r.bold=True;r.font.size=Pt(9.5)
    p=para(str(path));p.paragraph_format.space_after=Pt(5)
    for r in p.runs:r.font.name='Consolas';r.font.size=Pt(8)
    if extra:note(extra)

def build():
    plot_charts()
    para('医药产品组合分析报告','Title')
    para('市场表现 竞争动态与机器学习需求预测','Subtitle')
    para('IQVIA 商业数据科学家终面案例 中文核对版')
    para('原分析日期 2026 年 9 月 19 日    核对日期 2026 年 9 月 20 日')
    note('本报告为公开开发数据集的面试演示案例，不是 IQVIA 正式交付物。2024 年为原案例的预测年度，结果并非基于 2026 年新数据作出的当前预测。')
    head('执行摘要')
    para('核对结论：报告主要数值与现有代码结果一致，原代码重算也复现了现有输出；但若干验证方法和指标口径的文字描述需要修正。中文版保留原有数值与章节内容，并以“核对说明”标明修正，详细证据见下一页。')
    table(['SKU 数','通用名数','月度期数','2023 年销量'],[['1,730','20','60','122.2 百万']],size=11)
    para('20 个通用名的样本组合在 2023 年实现 122.2 百万包装当量销量，同比下降 1.3%，较 2021 年峰值下降 22.3%。组合规划不宜假设需求会自然回到此前增长轨道。按验证结果选择的阻尼 ETS 模型给出 2024 年基准预测 114.4 百万，同比下降 6.4%。')
    para('随机森林分类器在测试锚点 2022 年 7 月至 2023 年 6 月的合并观测中，前 20% 排名精确率为 58.3%，提升倍数为 2.33。六个月份额变化是最重要的预测信号；鉴于时间信息泄漏问题，该表现只能作为探索性证据。')
    para('普瑞巴林、乙酰半胱氨酸及对乙酰氨基酚的 2024 年预测增幅居前。布洛芬同时具备规模、正向份额动量和预测增长，是主要增长关注对象；前两项低份额产品应先验证需求来源，再考虑追加商业投入。')
    table(['行动','重点','依据'],[
        ['投入评估','布洛芬；验证普瑞巴林与乙酰半胱氨酸的需求驱动因素','份额动量与预测增长均为正'],
        ['防守','乙酰水杨酸、对乙酰氨基酚、双氯芬酸','规模较大且近期组合份额下滑'],
        ['供应规划','以阻尼 ETS 为案例基准，监测 -8.4% 测试偏差','关注低估需求的风险'],
        ['月度监测','份额动量、预测误差、断供迹象和行动分组','形成持续预警流程']], [2,9.1,6.1])
    note('数据 D1；分析代码 C1；结果 R1、R2、R5、R6。所有编号的完整绝对路径见附录 B。销量采用原分析的包装当量口径，不代表收入。')

    page('核对结果与必要修正')
    para('在独立输出目录重新运行原分析脚本，未修改原始数据、英文报告及原有分析结果。原始 CSV 的 SHA-256 与英文报告附录一致。')
    table(['核对对象','结果','证据'],[
        ['完整结果 JSON','一致','1,746 个数值、921 个其他值一致；最大数值差为 0'],
        ['7 个 CSV 输出','一致','逐表核对行列结构与全部内容通过'],
        ['报告 7 张主要结果表','一致','181 个所选数据单元格一致，包含数值及标识'],
        ['原始数据独立汇总','一致','1,730 个唯一 SKU；2019 至 2023 年年度总量一致'],
        ['关键正文数值与图表数据','一致','总量、增速、峰值差、模型指标和预测总量可追溯']], [4,2,11.2])
    head('核对说明',2)
    para('一是分类测试期。原摘要和复现附录的“2023 年留出集”不够准确。分类模型的测试锚点实际为 2022 年 7 月至 2023 年 6 月，其六个月后标签对应 2023 年 1 月至 12 月。需求预测模型才以 2023 年全年作为测试期。')
    para('二是时间信息泄漏。季节性强度以完整 2019 至 2023 年销量计算，再作为分类器及全局预测模型的静态特征；因此“所有 2023 年预测均未使用 2023 年结果”不适用于全局岭回归、全局提升树和包含它们的集成。季节性朴素法和阻尼 ETS 的单体预测不读取该静态特征，但整个候选模型比较流程仍应重新严格回测。')
    para('三是标签可得时间。分类训练锚点截至 2021 年 12 月，但标签最晚需要 2022 年 6 月数据；验证锚点截至 2022 年 6 月，标签最晚需要 2022 年 12 月数据。直接按锚点切分，不能保证在验证或测试起点已获得全部训练和选模标签。应按标签成熟日期设置时间间隔并采用滚动验证。')
    para('四是指标口径。代码将整个测试期 240 条观测合并排名，选择前 48 条计算 Precision@20%；这不等于逐月筛选 20 个通用名中的前 4 个。原文 PR-AUC 实际使用自定义平均精确率 AP；其并列分数处理依赖排序，不能直接视作所有标准库 PR-AUC 实现的同一指标。')
    para('五是区间解释。每个预测月份的范围取自 2022 年同一预测步长下 20 个通用名对数残差的 10% 与 90% 分位数，再逐月相加。该方法没有校验覆盖率；年度上下界之和不能声称具有已验证的年度 80% 覆盖概率。')
    para('其他表述已澄清：将“-6.4% below”改为“下降 6.4%”；保留样本中的 ATC 分类而不将其解释为外部标准校验；“市场带动增长”是代码规则名称，并不证明增长原因。')
    note('证据 C1 第 563 至 629、733 至 746、787 至 846 行；复核代码 C3；核对记录 V1。数值可复现不等于模型已通过严格前瞻验证。')

    page('1 商业问题与分析设计')
    data_page_start=len(DOC.paragraphs)-1
    para('本案例研究需求增长与下降、样本内份额得失、未来六至十二个月值得关注的通用名，以及如何结合预测和分群配置商业分析资源。分类目标为未来六个月份额表现，需求预测为未来十二个月。')
    para('分析流程依次为 SKU 销量整理、文本特征提取、市场定义、竞争分析、组合分群、预警分类、需求预测和透明行动规则。原代码使用时间顺序切分，没有随机拆分训练与测试集；但其严格时间有效性仍受前页说明的特征与标签问题限制。')
    head('1.1 数据集与工程处理',2)
    table(['要素','实现'],[['粒度','每个 Morion ID 对应一行药品 SKU'],['覆盖','1,730 个 SKU；20 个 INN；2019 年 1 月至 2023 年 12 月'],['销量','60 个宽表月度字段，转成长表并汇总至 INN 与月份'],['层级','ATC1 至 ATC5、通用名、匿名品牌、SKU'],['文本特征','从 Description 提取规格数值、单位和包装数量；从 NFC (1) 简化剂型']], [3,14.2])
    head('1.2 数据质量',2)
    table(['检查','结果'],[['完整性','0 个销量缺失值；0 个负销量值；SKU ID 无重复'],['零销量','占 SKU 月度观测的 32.0%，保留为可得性或间歇需求信号'],['文本提取','规格提取率 98.6%；包装数量提取率 97.3%'],['潜在断供','305 个前后两月均为正而当月为零的间断点；保留观测'],['潜在尖峰','161 个局部尖峰；保留用于汇总需求及敏感性评估']], [3,14.2])
    para('核对说明：上述断供和尖峰是启发式标记，不能直接认定为真实缺货或错误数据。尖峰实际判断为当月销量大于 max（局部正值中位数的 5 倍，中位数 + 50）。原代码只汇总这些标记的数量，没有将逐 SKU 月份的标记明细导出。')
    head('1.3 市场定义',2)
    para('优先使用至少包含两个样本通用名的 ATC4 分组，否则回退到符合条件的 ATC3 分组。仍无样本同类产品时，不报告竞争份额。20 个通用名中有 12 个至少出现在一个可用同类分组，其余使用明确标识的组合份额。')
    para('该文件是选取的 20 个通用名样本，不是完整市场普查。份额仅代表样本；数据不足以支持收入、价格弹性、推广回报、患者或医生定向、准入、利润及停产品决策。')
    note('数据 D1；代码 C1 第 308 至 341、764 至 823、879 至 889 行；结果 R1、R4。')
    for p in DOC.paragraphs[data_page_start:]:
        if not p.text.strip():continue
        p.paragraph_format.line_spacing=Pt(20 if p.style.name.startswith('Heading') else 15)
        p.paragraph_format.space_before=Pt(4 if p.style.name.startswith('Heading') else 0)
        p.paragraph_format.space_after=Pt(4)
        snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');p._p.get_or_add_pPr().append(snap)

    page('2 市场表现与竞争动态')
    figure('annual.png','图 1  2019 至 2023 年组合年度销量')
    para('组合需求在 2020 年增长 8.2%，2021 年增长 1.4%，随后在 2022 年下降 21.3%，2023 年下降 1.3%。2023 年看似企稳的总量掩盖了通用名之间的明显分化。')
    head('2.1 通用名表现',2)
    table(['通用名','2023 销量\n百万','同比','组合份额','份额变化\n百分点'],[[M[r.INN],million(r.sales_2023),pct(r.yoy_growth_2023),pct(r.portfolio_share_2023),pp(r.portfolio_share_change_2023)] for _,r in K.nlargest(6,'sales_2023').iterrows()],[4.5,3.1,2.7,3.2,3.7])
    para('乙酰水杨酸以 22.4 百万销量位居第一，但组合份额减少 0.5 个百分点。布洛芬增长 8.4%，组合份额增加 1.0 个百分点，是高销量产品中动量较强的品种。对乙酰氨基酚仍居第三，但销量下降 12.6%，组合份额减少 1.3 个百分点。')
    note('图表及计算依据 D1、R1 portfolio_annual、R2、R3；代码 C1 第 774 至 789 行。图中单位为百万包装当量，正文保留原报告的一位小数精度。')

    page('2.2 样本同类市场动态')
    peers=pd.DataFrame(R['peer_analysis']);peers=peers[peers.market_name.str.contains('HMG|M01A |N02B ',regex=True)].sort_values(['market_name','share_change_pp'],ascending=[True,False])
    market=lambda x:'C10A A 他汀类' if x.startswith('C10') else ('M01A 全身用非甾体抗炎药' if x.startswith('M01') else 'N02B 其他镇痛及解热药')
    table(['通用名','样本市场','市场增长','份额变化\n百分点','份额效应\n包装当量'],[[M[r.INN],market(r.market_name),pct(r.market_growth_2023),pp(r.share_change_pp),f'{r.share_effect:,.0f}'] for _,r in peers.iterrows()],[3.5,4.5,2.5,2.9,3.8],8.7)
    para('样本他汀类市场增长 19.0%，其中瑞舒伐他汀份额增加 3.0 个百分点，阿托伐他汀相应减少 3.0 个百分点。市场扩张与产品结构变化均对销量结果有贡献。')
    para('在全身用非甾体抗炎药分组中，布洛芬份额增加 2.0 个百分点，份额效应约为正 51.9 万包装当量；双氯芬酸份额减少 3.1 个百分点，份额效应约为负 79.5 万包装当量。')
    para('在原始数据的其他镇痛及解热药分组中，市场整体下降 4.6%，普瑞巴林份额仍增加 3.7 个百分点，份额效应约为正 99.7 万；对乙酰氨基酚份额减少 4.1 个百分点，份额效应约为负 109.6 万包装当量。')
    head('增长拆解的商业含义',2)
    para('市场扩张效应等于按上期份额分配的当期市场销量减去上期产品销量；份额效应等于当期实际销量减去按上期份额分配的当期市场销量。两者之和等于销量变动。该拆解区分了类别整体变化与相对份额变化，但不能据此认定具体商业行为的因果效果。')
    para('核对说明：原始文件中的 ATC 字段是本节分组的直接依据，例如普瑞巴林归入上述 N02B 样本分组。此次核对确认代码忠实使用原始标签，未将该标签当作已按现行外部分类标准校验的结论。每个市场只包含该分组内的 SKU，不能与某通用名全部剂型的总销量直接混用。')
    note('数据 D1；结果 R7 market_peer_analysis.csv；代码 C1 第 790 至 809 行。此表保留原报告选择的三个样本市场，完整输出另含其他分组。')

    page('3 产品组合分群')
    para('将通用名层面的规模、增长、份额动量、需求波动、季节性和产品复杂度等 14 项特征标准化，比较 K 为 3、4、5 的 K 均值聚类。K 为 4 时轮廓系数最高，为 0.232；其余分别为 0.208 和 0.211，因此原代码采用四类描述性分群。')
    figure('clusters.png','图 2  四类产品组合的主成分投影')
    table(['类型','INN 数','2023 销量\n百万','平均同比','平均份额变化\n百分点','平均预测\n增长率'],[[CL[r['cluster_label']],r['INN_count'],million(r['sales_2023']),pct(r['avg_yoy_growth']),pp(r['avg_share_change']),pct(r['avg_forecast_growth'])] for r in R['cluster_summary']],[3.4,1.6,3,2.8,3.4,3],8.5)
    para('成长与新兴型包含普瑞巴林和瑞舒伐他汀，近期动量较强。成熟领先型包含规模较大的成熟产品。风险与下滑型的平均月度变异系数最高，为 0.552，且历史增长为负。分群只用于描述，商业行动建议另由预测和份额规则决定。')
    para('核对说明：表中“平均”是组内通用名指标的简单算术平均，并非销量加权增长率；“利基与波动型”是原代码的启发式名称，该组实际平均变异系数最低，为 0.147，不应仅凭名称判断其波动程度。')
    note('结果 R1 cluster_validation、cluster_summary；R2；代码 C1 第 750 至 761、858 至 872 行。PCA 仅用于可视化。')

    page('4 商业机会分类模型')
    para('目标“未来份额赢家”定义为：在某个锚点月份，未来六个月组合份额增量排名前 25% 的通用名。每月 20 个通用名取前 5 个，基础赢家比例为 25%。由于多数样本细分市场缺少稳定的同类产品集合，分类标签采用组合份额。')
    table(['用途','锚点月份','观测数','六个月后标签时间'],[['训练','2020 年 1 月至 2021 年 12 月','480','2020 年 7 月至 2022 年 6 月'],['验证与选模','2022 年 1 月至 2022 年 6 月','120','2022 年 7 月至 2022 年 12 月'],['测试','2022 年 7 月至 2023 年 6 月','240','2023 年 1 月至 2023 年 12 月']],[2.8,5.8,1.8,6.8],8.6)
    para('选模优先比较验证期 Precision@20%，并依次以 AP 和 ROC-AUC 打破并列。随机森林因此胜出，尽管提升树的验证 AP 更高。代码所用随机森林为 100 棵树、深度 3、最小叶节点样本数 20。')
    table(['模型','验证 AP','测试 AP','测试 P@20%','测试提升倍数'],[[MODEL[r['model']],pct(r['validation_pr_auc']),pct(r['test_pr_auc']),pct(r['test_precision_top20']),f"{r['test_lift_top20']:.2f}"] for r in R['classification_results']],[4.8,3.1,3.1,3.2,3],9)
    figure('lift.png','图 3  整个分类测试期合并排名的前 20% 提升倍数',15.0)
    para('随机森林测试 ROC-AUC 为 75.2%，AP 为 61.3%，前 20% 精确率为 58.3%，提升倍数为 2.33。L2 逻辑回归的测试提升倍数最高，但不应据此在同一测试集上重新选择模型。上述结果包含前述时间信息泄漏风险，尚不能据此确认实际投用时的排名收益。')
    note('结果 R5；代码 C1 第 583 至 629、824 至 846 行。AP 为原表 PR-AUC 列的实际计算口径；P@20% 为合并排名而非逐月排名。')

    page('4.1 模型解释与使用边界')
    para('原代码在测试集上逐个打乱特征，每项重复 15 次，计算平均 AP 下降量。六个月份额变化是最强信号；规格种类数、六个月增长和三个月份额变化也有贡献。这些是预测关联，不是能保证带来增长的干预因素。')
    ft={'share_change6':'六个月组合份额变化','log_strength_count':'规格种类数的对数变换','growth6':'六个月对数销量变化','share_change3':'三个月组合份额变化','relative_growth12':'十二个月相对对数增长'}
    table(['特征中文名称','代码字段','打乱后 AP 下降'],[[ft[r['feature']],r['feature'],f"{r['pr_auc_drop']:.3f}"] for r in R['classification_importance'][:5]],[6.5,6.3,4.4])
    head('实际部署前需要的修复',2)
    para('所有随时间变化的特征都应只使用锚点可见数据。季节性强度应按历史窗口重算；规格、包装及 SKU 数量还需要带时间戳的产品主数据，确认在对应历史时点已经可得。')
    para('在每次训练及选择模型时，应排除尚未经过完整六个月、标签尚不成熟的锚点。随后采用多个滚动时间起点评估模型，并分别报告每月前 4 个通用名的精确率、提升倍数及跨期稳定性。')
    para('核对说明：原代码完成了历史分类评估和特征重要性计算，没有保存可直接加载的分类模型，也没有导出 2023 年末或当前月份的逐通用名赢家概率。“月度评分”属于原报告建议的后续运行流程，而非已经实现的生产功能。')
    head('结果的商业用法',2)
    para('可以据此把六个月份额动量纳入月度观察清单，并据具体品种核查渠道、供应、剂型和竞争变化。投入决策应在确认驱动因素后作出，不能仅凭模型概率或特征重要性推断干预回报。')
    note('结果 R1 classification_importance；代码 C1 第 617 至 629 行；特征构建第 824 至 846 行。严格回测修复属于下一步建议，本次保留原模型输出用于核对。')

    page('5 需求预测')
    para('在通用名与月份层面比较季节性朴素法、阻尼 ETS、带日历及滞后特征的全局岭回归、浅层全局提升树和验证误差加权集成。2019 至 2021 年数据用于预测 2022 年并选择模型，再用截至 2022 年的数据拟合后预测 2023 年。')
    table(['模型','验证 WAPE','测试 WAPE','测试 MASE','测试偏差'],[[MODEL[r['model']],pct(r['validation_WAPE']),pct(r['test_WAPE']),f"{r['test_MASE']:.2f}",pct(r['test_Bias'])] for r in R['forecast_results']],[4.8,3.1,3.1,3.1,3.1],9)
    para('阻尼 ETS 在 2022 年验证期以 28.4% WAPE 胜出。2023 年测试 WAPE 为 23.1%，MASE 为 0.67，偏差为 -8.4%。负偏差表示整体低估需求，应在供应规划中持续监测。')
    figure('forecast.png','图 4  2023 年组合需求与阻尼 ETS 及季节性基准预测')
    para('验证加权集成的 2023 年 WAPE 为 16.3%，低于阻尼 ETS，但它没有在验证期胜出，且包含使用全时期静态特征的全局模型。该结果应作为修复后前瞻回测的候选假设，不应被当作可直接替换基准模型的依据。')
    para('核对说明：表内误差是在 20 个通用名与 12 个月构成的 240 个单元上计算，图则是各月组合合计。图形吻合程度不能直接替代通用名层面的误差评价。MASE 使用所有通用名训练期季节性绝对差的合并平均值作为分母。')
    note('结果 R6、R1 forecast_2023_portfolio；代码 C1 第 631 至 747 行。原报告“所有预测均未使用 2023 年结果”的说法已按核对结果修正。')

    page('5.1 2024 年基准预测')
    para('将验证选出的阻尼 ETS 用完整 2019 至 2023 年数据重新拟合，并自下而上加总通用名预测，得到 2024 年组合销量 114.4 百万包装当量，较 2023 年下降 6.4%。未四舍五入的总量为 114,362,650.51，同比为 -6.4269%。')
    top=K.nlargest(8,'forecast_growth_2024')
    table(['通用名','2023 销量\n百万','2024 预测\n百万','增长率','行动分组'],[[M[r.INN],million(r.sales_2023),million(r.forecast_2024),pct(r.forecast_growth_2024),ACT[r.action_segment]] for _,r in top.iterrows()],[4.5,3,3,2.8,3.9],8.8)
    figure('matrix.png','图 5  预测增长 组合份额动量与产品规模',15.6)
    para('布洛芬是规模较大的增长领先品种；普瑞巴林和乙酰半胱氨酸为新兴机会。对乙酰氨基酚与乙酰水杨酸虽预测增长，但份额动量为负，应同时开展防守分析。左氧氟沙星兼有负份额动量及最弱的预测增长，是最明确的风险信号。')
    para('核对说明：原文“80% 范围”是基于 2022 年对数残差分位数的经验范围，未验证实际覆盖率，也未纳入上市、短缺、价格或政策情景。“市场带动增长”是规则名称：它适用于预测增长但份额未增加且份额低于中位数的产品，并未识别真正的市场驱动原因。')
    note('结果 R2、R8；代码 C1 第 733 至 746、847 至 857、874 至 877 行。气泡图采用组合份额；绿色表示预测增长，红色表示预测下降。')

    page('6 商业建议与月度运行机制')
    table(['建议','范围','行动'],[
        ['优先关注增长领先品种','布洛芬','保持商业支持与供应保障，核查带动增长的规格和剂型'],
        ['培育新兴机会','普瑞巴林、乙酰半胱氨酸','在增量投入前验证需求、分销覆盖和制剂结构变化'],
        ['防守高销量产品','乙酰水杨酸、对乙酰氨基酚、双氯芬酸','诊断样本份额流失、竞争替代及剂型层面的转移'],
        ['管理预测风险','整个组合','阻尼 ETS 保留为案例基准，月度监测偏差，并在修复后前瞻比较集成模型'],
        ['建立预警流程','全部通用名','按月更新增长、样本份额、份额变化、预测误差和行动分组']],[3.5,4.2,9.5],9.2)
    head('6.1 月度决策节奏',2)
    table(['步骤','输出'],[['数据刷新','加载最新月份，重新运行质量标记和特征提取'],['表现复盘','更新销量增长、同类样本份额、份额变化及增长拆解'],['模型监测','修复并实现月度评分后，跟踪未来赢家概率及各通用名预测误差'],['行动复盘','对达到透明增长、份额及预测阈值的品种开展进一步分析'],['学习反馈','记录商业假设与结果，为后续分析补充因果背景']],[3.5,13.7],9.5)
    para('运行原则：不要对所有通用名采用相同的预测和商业处理方式。稳定的成熟需求、波动需求与间歇需求需要不同的监测阈值和规划方法。当前分析尚不具备完整自动刷新、模型保存和前瞻评分流程。')
    head('行动规则的精确定义',2)
    para('“高份额”为 2023 年组合份额不低于 20 个通用名的份额中位数。若预测增长为正且份额增加，高份额品种归为增长领先，低份额品种归为新兴机会。高份额且份额未增加者优先归为防守；其余预测与份额均未改善者为风险；余下预测为正者标记市场带动增长，否则监测。')
    note('行动规则 R2 action_segment；代码 C1 第 850 至 857 行。商业建议属于对分析结果的解释，并非已验证的投入回报或因果结论。')

    page('7 局限性与下一步')
    table(['局限','对决策的影响'],[['样本完整性','20 个通用名并不构成完整市场；样本份额不是外部市场份额'],['商业驱动变量','缺少价格、推广、准入、支付方、患者、医生、供应政策及竞争进入数据'],['结果解释','预测重要性不证明因果关系，也不能指定能保证增长的干预'],['时间长度','只有 60 个月观测；结构变化可能主导年度回测表现'],['验证有效性','全时期季节性特征与六个月标签跨界削弱严格样本外结论'],['区间与情景','经验范围缺少覆盖率检验；未纳入上市、短缺与政策情景'],['盈利能力','销量不能支持利润、收入或停产品决策']],[3.5,13.7],9.5)
    para('下一阶段应取得覆盖更完整竞争产品的市场数据，重做 ATC4 份额建模，并核验原始分类标签。若涉及盈利或销售队伍决策，还需加入价格、净销售额、成本、推广触达、支付方准入及患者或医生数据。')
    para('方法层面应首先修复时间可得性：按历史锚点构造特征，按六个月标签成熟日期设置间隔，再扩展滚动起点回测。标准化 AP 及并列分数处理，增加逐月排名指标、通用名分层误差和预测区间覆盖率评估。')
    para('对 2024 年预测应保持原案例语境。没有 2024 年及之后的实际销量，无法在本次核对中评价该年度预测的实际准确性，也不能将其转述为当前经营预测。')
    head('本次核对边界',2)
    para('本次完成了原始 CSV 指纹核验、原代码独立重跑、现有 JSON 与 7 个 CSV 全量比对、报告主要结果表及关键正文数值核对，并在中文报告中补充口径说明。没有调整原模型或重估一套“修复后指标”，以避免将模型修订与原报告一致性核对混为一项结果。')
    para('中文图表由核对通过的数据重新绘制；药品拉丁名称保留在附录 C，便于与代码字段准确对应。数据路径、代码路径、结果路径与复现命令列于附录 B。')

    page('附录 A 技术指标定义')
    table(['指标','定义及本案例口径'],[
        ['同比增长','当期销量 ÷ 上年同期销量 − 1；年销量比较使用当年与上年合计。分母为零时无定义。'],
        ['复合年增长率','（2023 年销量 ÷ 2019 年销量）的 1/4 次方 − 1。'],
        ['样本市场份额','通用名在合格 ATC4 或回退 ATC3 分组中的销量 ÷ 该分组样本总销量。'],
        ['组合份额','通用名全部样本销量 ÷ 全部 20 个通用名的样本总销量。'],
        ['份额变化','当期份额减上期份额。代码中以比例保存，报告乘 100 后以百分点展示。'],
        ['份额效应','当期实际销量 − 当期样本市场总销量 × 上期份额。'],
        ['季节性强度','先去除线性趋势，以各月份平均值构造季节项；1 − 残差方差 ÷（季节项加残差）的方差，截断至 0 至 1。'],
        ['未来份额赢家','每个锚点下，未来六个月组合份额增量前 25% 的通用名，每月 5 个。'],
        ['AP','按预测分数降序，对每个正例位置的累计精确率求平均；原代码列名为 pr_auc。并列分数按排序结果处理。'],
        ['Precision@20%','该评估分区全部观测合并排名后，前 20% 观测中的赢家比例。测试期为前 48 条。'],
        ['Lift@20%','Precision@20% ÷ 全部观测赢家比例；本例基础比例为 25%。'],
        ['WAPE','绝对预测误差之和 ÷ 实际销量之和；按通用名与月份单元合并计算。'],
        ['MASE','预测绝对误差平均值 ÷ 训练数据 12 个月季节性差分绝对值的合并平均值。'],
        ['偏差 Bias','预测值减实际值的差额之和 ÷ 实际销量之和。负值表示总体低估。'],
        ['经验范围','每一预测月份使用 2022 年同一步长跨通用名对数残差的 10% 与 90% 分位数；未保证覆盖率。']],[3.7,13.5],9.0)
    note('计算定义以 C1 的实际实现为准。代码使用 NumPy 自定义分类、树模型、预测及聚类实现；不能假定与同名标准库模型在算法细节上完全等价。')

    page('附录 B 数据与代码位置')
    para('以下为本次核对所使用文件的完整绝对路径。D 表示输入数据，C 表示代码，R 表示原有结果，V 表示本次复核记录。')
    location('D1','原始数据',ROOT/'data/teva_sales_small.csv','宽表 1,730 行、70 列，其中 60 列为月度销量。SHA-256 如下。')
    note(V['data_sha256'])
    location('C1','核心分析代码',ROOT/'analysis/iqvia_case_analysis.py','数据与份额分析 764–823 行；分类 563–629、824–844 行；预测 631–747、845–857 行；聚类 858–872 行；输出 899–933 行。')
    location('C2','英文报告生成代码',ROOT/'analysis/build_report.py')
    location('C3','本次数值复核代码',ROOT/'analysis/verify_report_20260920.py')
    location('C4','中文报告与中文图表生成代码',ROOT/'analysis/build_chinese_report_20260920.py')
    location('R1','完整分析结果',BASE/'analysis/analysis_results.json')
    location('R2','通用名 KPI 与预测',BASE/'analysis/inn_kpi.csv')
    location('R3','通用名月度销量',BASE/'analysis/inn_monthly.csv')
    location('R4','SKU 特征与年度销量',BASE/'analysis/sku_features.csv')
    note('本分析输入为 D1。工作区 data/raw 下的其他原始数据和 R/01_download.R 未被 C1 引用，不属于这份 Teva 案例报告的数据链路。')

    page('附录 B 结果文件与复现步骤')
    location('R5','分类模型指标',BASE/'analysis/classification_model_results.csv')
    location('R6','预测模型指标',BASE/'analysis/forecast_model_results.csv')
    location('R7','同类市场分析',BASE/'analysis/market_peer_analysis.csv')
    location('R8','2024 年月度预测与经验上下界',BASE/'analysis/forecast_2024_monthly.csv')
    location('V1','本次机器可读核对记录',QA/'verification_summary.json')
    location('V2','独立重跑结果目录',QA/'recomputed')
    head('复现命令',2)
    para('以项目根目录 E:\\iqvia商业分析实例 为工作目录，按以下顺序执行。python 指向支持 NumPy、pandas、Pillow 及 python-docx 的 Python 解释器。')
    for cmd in ['python analysis/iqvia_case_analysis.py --input data/teva_sales_small.csv --output outputs/iqvia_teva_case_20260919/verification_20260920/recomputed','python analysis/verify_report_20260920.py','python analysis/build_chinese_report_20260920.py']:
        p=para(cmd)
        for r in p.runs:r.font.name='Consolas';r.font.size=Pt(8)
    para('随机种子为 20260919。本次重算使用 NumPy 2.3.5 和 pandas 3.0.1。原始数据、原代码和英文报告的 SHA-256 保存在 V1，便于后续确认版本。预测采用通用名预测相加的自下而上汇总。')
    note('原报告来源链接：Teva 开发数据项目 https://github.com/Marchev-Science/case-forecasting-pharmacutical-demand ；WHO ATC 分类 https://www.who.int/tools/atc-ddd-toolkit/atc-classification ；IQVIA 品牌与组合策略 https://www.iqvia.com/solutions/commercialization/brand-and-portfolio-strategy 。上述链接沿用原报告，数值核对以本地 D1 和代码结果为准。')

    page('附录 C 通用名与代码标识对应')
    para('正文使用中文通用名便于阅读。下表保留数据中的拉丁名称及标记，供检索 CSV 和代码输出时使用。乙酰水杨酸后的星号沿用原始字段。')
    table(['中文通用名','原始 INN 字段','原代码行动分组'],[[M[r.INN],r.INN,ACT[r.action_segment]] for _,r in K.iterrows()],[4.2,9,4],9)
    note('来源 R2。行动分组基于 2023 年份额与原案例 2024 年预测，不代表当前市场判断。')
    DOC.core_properties.title='医药产品组合分析报告 中文核对版'
    DOC.core_properties.subject='原报告与代码结果一致性核对 中文翻译 数据与代码溯源'
    DOC.core_properties.author=''
    output=BASE/'IQVIA_Pharma_Portfolio_Analytics_Report_中文核对版.docx'
    DOC.save(output)
    print(str(output))

if __name__=='__main__':build()
